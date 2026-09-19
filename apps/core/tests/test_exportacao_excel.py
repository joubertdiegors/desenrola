"""
Exportação para Excel -- Usuários e Cartas do Backoffice (Rodada 20).

O QUE ESTA SUÍTE EXISTE PARA IMPEDIR
------------------------------------
1. **Uma brecha pela URL.** A exportação tem as MESMAS portas da
   listagem: anônimo vai para o login, sem a permissão é 403, GET é 405,
   e sem o token CSRF o POST é recusado;
2. **Uma planilha que não é planilha.** A resposta é um .xlsx de verdade
   (um zip), aberto aqui com openpyxl -- não um CSV renomeado;
3. **Coluna a mais.** "Tabela completa" traz as colunas da tela, sem
   Ações; "Contatos", SÓ nome, e-mail e telefone. Nada de senha,
   documento, IDs, permissões;
4. **O telefone virar número.** "+32 470 00 00 00" e "0470..." chegam
   como TEXTO, iguais ao cadastro;
5. **Uma exportação que discorda da tela.** A busca e os filtros em
   vigor valem; a paginação NÃO: sai o resultado inteiro, não a página;
6. **Um nome que vira fórmula.** "=..." continua texto.
"""

import datetime
import io
import re

import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.letters import services
from apps.letters.models import Letter

pytestmark = pytest.mark.django_db

USUARIOS = reverse("backoffice:users")
CARTAS = reverse("backoffice:letters")
EXPORTAR_USUARIOS = reverse("backoffice:users_export")
EXPORTAR_CARTAS = reverse("backoffice:letters_export")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CABECALHO_USUARIOS = [
    "Nome", "E-mail", "Telefone", "Endereço", "Cidade", "Cadastro", "Último acesso",
    "Cartas", "Status",
]
CABECALHO_CARTAS = [
    "Referência", "Criada por", "Para", "Idioma", "Data de ida", "Data de volta", "Status",
    "Criada em", "Finalizada em",
]
CABECALHO_CONTATOS = ["Nome", "E-mail", "Telefone"]

BACKOFFICE = ("core", "access_backoffice")
GERENCIA = ("accounts", "manage_users")
SUPERVISAO = ("letters", "view_all_letters")


@pytest.fixture(autouse=True)
def _senha_rapida(settings):
    """
    Um hash de senha rápido, SÓ nesta suíte: ela cria centenas de contas,
    e o PBKDF2 de produção levaria minutos só nisso. A senha não é o
    assunto aqui -- e o teste de dado sensível continua conferindo que o
    hash não sai na planilha.
    """
    settings.PASSWORD_HASHERS = ["django.contrib.auth.hashers.MD5PasswordHasher"]


def _cliente(*permissoes, email="operadora-excel@mail.com", csrf=False):
    pessoa = get_user_model().objects.create_user(
        email=email, password="x", full_name="Operadora Excel"
    )
    for app_label, codename in permissoes:
        pessoa.user_permissions.add(
            Permission.objects.get(content_type__app_label=app_label, codename=codename)
        )
    c = Client(enforce_csrf_checks=csrf)
    c.force_login(pessoa)
    return c


@pytest.fixture
def gerente():
    return _cliente(BACKOFFICE, GERENCIA)


@pytest.fixture
def supervisora():
    return _cliente(BACKOFFICE, SUPERVISAO)


def _planilha(resposta):
    """A resposta tem de ser um .xlsx DE VERDADE -- e é aberta com openpyxl."""
    assert resposta.status_code == 200, resposta.status_code
    assert resposta["Content-Type"] == XLSX
    assert resposta.content[:2] == b"PK"  # um zip, como todo .xlsx; CSV não é
    return load_workbook(io.BytesIO(resposta.content)).active


def _linhas(folha):
    return [list(linha) for linha in folha.iter_rows(values_only=True)]


def _dados(folha):
    return _linhas(folha)[1:]


def _textos(folha):
    return [str(valor) for linha in _linhas(folha) for valor in linha if valor is not None]


def _data(valor):
    """Como o openpyxl devolve uma célula de data: datetime à meia-noite."""
    assert isinstance(valor, datetime.datetime), valor
    return valor.date()


def _hoje():
    return timezone.localdate().isoformat()


def _formulario_de_exportacao(html):
    """O formulário do "Exportar Excel ▾", recortado da página."""
    padrao = r'<details class="mod-menu mod-menu-exportar[^"]*">.*?</details>'
    achado = re.search(padrao, html, re.S)
    assert achado, "o controle de exportação não está na página"
    return achado.group(0)


def _escondidos(formulario):
    """Os filtros que o formulário leva escondidos (fora o token CSRF)."""
    campos = dict(re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)">', formulario))
    campos.pop("csrfmiddlewaretoken", None)
    return campos


def _colunas_da_tabela(html):
    """Os títulos das colunas da tabela da tela, na ordem."""
    cabecalho = re.search(r'<div class="mod-cabecalho-tabela" role="row">(.*?)</div>', html, re.S)
    assert cabecalho, "a tabela não está na página"
    return re.findall(r'<span role="columnheader"[^>]*>([^<]+)</span>', cabecalho.group(1))


# ===========================================================================
# USUÁRIOS
# ===========================================================================


def _pessoa(nome, email, **dados):
    return get_user_model().objects.create_user(
        email=email, password="senha-que-nao-sai", full_name=nome, **dados
    )


@pytest.fixture
def pessoas():
    completa = _pessoa(
        "Ana Completa", "ana@exemplo.be",
        phone="+32 470 12 34 56", address_line1="Rue Haute 10", address_line2="Boîte 2",
        city="Bruxelles", document_number="DOC-SECRETO-123",
        birth_date=datetime.date(1980, 1, 2),
    )
    get_user_model().objects.filter(pk=completa.pk).update(
        last_login=timezone.make_aware(datetime.datetime(2026, 9, 1, 10, 30))
    )
    vazia = _pessoa("Bruno Vazio", "bruno@exemplo.be")
    inativa = _pessoa("Carla Inativa", "carla@exemplo.be", phone="0470 99 88 77", is_active=False)
    return {
        "completa": get_user_model().objects.get(pk=completa.pk),
        "vazia": vazia,
        "inativa": inativa,
    }


def _exportar_usuarios(cliente, tipo, **filtros):
    return cliente.post(EXPORTAR_USUARIOS, {"tipo": tipo, **filtros})


def _linha_de(folha, email, coluna_do_email=1):
    for linha in _dados(folha):
        if linha[coluna_do_email] == email:
            return linha
    raise AssertionError(f"{email} não está na planilha")


class TestUsuariosPortas:
    def test_anonimo_vai_para_o_login(self, client):
        resposta = client.post(EXPORTAR_USUARIOS, {"tipo": "completa"})

        assert resposta.status_code == 302
        assert "/accounts/login/" in resposta.url

    def test_get_nao_exporta(self, gerente):
        assert gerente.get(EXPORTAR_USUARIOS, {"tipo": "completa"}).status_code == 405

    @pytest.mark.parametrize(
        "permissoes",
        [
            (),
            (BACKOFFICE,),
            (GERENCIA,),
            (BACKOFFICE, SUPERVISAO),
        ],
        ids=["nenhuma", "so-backoffice", "so-gerencia-sem-backoffice", "supervisao-de-cartas"],
    )
    def test_sem_a_permissao_da_listagem_e_403(self, permissoes):
        cliente = _cliente(*permissoes)

        for tipo in ("completa", "contatos"):
            assert cliente.post(EXPORTAR_USUARIOS, {"tipo": tipo}).status_code == 403

    def test_superusuario_exporta(self):
        dono = get_user_model().objects.create_superuser(email="dono@mail.com", password="x")
        c = Client()
        c.force_login(dono)

        assert _planilha(c.post(EXPORTAR_USUARIOS, {"tipo": "completa"}))

    def test_sem_o_token_csrf_e_recusado(self):
        cliente = _cliente(BACKOFFICE, GERENCIA, csrf=True)

        assert cliente.post(EXPORTAR_USUARIOS, {"tipo": "completa"}).status_code == 403

    @pytest.mark.parametrize("tipo", ["", "tudo", "senhas", "../completa"])
    def test_tipo_desconhecido_e_400(self, gerente, tipo):
        assert gerente.post(EXPORTAR_USUARIOS, {"tipo": tipo}).status_code == 400


class TestUsuariosTabelaCompleta:
    def test_e_um_xlsx_com_o_nome_do_dia(self, gerente, pessoas):
        resposta = _exportar_usuarios(gerente, "completa")

        _planilha(resposta)
        assert resposta["Content-Disposition"] == (
            f'attachment; filename="usuarios_tabela_completa_{_hoje()}.xlsx"'
        )

    def test_os_cabecalhos_sao_as_colunas_da_tela_sem_acoes(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "completa"))

        assert _linhas(folha)[0] == CABECALHO_USUARIOS
        assert "Ações" not in _linhas(folha)[0]

    def test_uma_linha_por_pessoa(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "completa"))

        assert len(_dados(folha)) == get_user_model().objects.count()

    def test_a_linha_traz_o_que_a_tabela_mostra(self, gerente, pessoas):
        ana = pessoas["completa"]

        linha = _linha_de(_planilha(_exportar_usuarios(gerente, "completa")), ana.email)

        assert linha[:5] == [
            "Ana Completa", "ana@exemplo.be", "+32 470 12 34 56", "Rue Haute 10, Boîte 2",
            "Bruxelles",
        ]
        assert _data(linha[5]) == timezone.localtime(ana.date_joined).date()
        assert _data(linha[6]) == datetime.date(2026, 9, 1)
        assert linha[7] == 0
        assert linha[8] == "Ativo"

    def test_a_quantidade_de_cartas_e_um_numero(self, gerente, pessoas, modelos_oficiais_prontos):
        for idioma in ("fr", "pt"):
            services.start_draft(pessoas["completa"], idioma)

        linha = _linha_de(_planilha(_exportar_usuarios(gerente, "completa")), "ana@exemplo.be")

        assert linha[7] == 2

    def test_inativo(self, gerente, pessoas):
        linha = _linha_de(_planilha(_exportar_usuarios(gerente, "completa")), "carla@exemplo.be")

        assert linha[8] == "Inativo"

    def test_o_que_falta_fica_vazio(self, gerente, pessoas):
        linha = _linha_de(_planilha(_exportar_usuarios(gerente, "completa")), "bruno@exemplo.be")

        # Telefone, endereço, cidade e último acesso: vazios -- nem "—",
        # nem "Nunca", que são da tela.
        assert linha[2] is None
        assert linha[3] is None
        assert linha[4] is None
        assert linha[6] is None

    def test_nenhum_dado_sensivel(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "completa"))
        textos = " | ".join(_textos(folha))
        ana = pessoas["completa"]

        assert ana.password not in textos
        assert "senha-que-nao-sai" not in textos
        assert "DOC-SECRETO-123" not in textos
        assert "1980" not in textos
        for proibido in ("Senha", "Documento", "Permiss", "ID", "Superusu"):
            assert not any(proibido in titulo for titulo in _linhas(folha)[0])


class TestUsuariosContatos:
    def test_so_nome_email_e_telefone(self, gerente, pessoas):
        resposta = _exportar_usuarios(gerente, "contatos")
        folha = _planilha(resposta)

        assert resposta["Content-Disposition"] == (
            f'attachment; filename="usuarios_contatos_{_hoje()}.xlsx"'
        )
        assert _linhas(folha)[0] == CABECALHO_CONTATOS
        assert folha.max_column == 3
        assert _linha_de(folha, "ana@exemplo.be") == [
            "Ana Completa", "ana@exemplo.be", "+32 470 12 34 56",
        ]
        assert _linha_de(folha, "bruno@exemplo.be") == ["Bruno Vazio", "bruno@exemplo.be", None]

    def test_nada_alem_do_contato(self, gerente, pessoas):
        textos = " | ".join(_textos(_planilha(_exportar_usuarios(gerente, "contatos"))))

        for fora in ("Rue Haute", "Bruxelles", "Ativo", "Inativo", "DOC-SECRETO-123",
                     pessoas["completa"].password):
            assert fora not in textos

    def test_uma_linha_por_pessoa(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "contatos"))

        assert len(_dados(folha)) == get_user_model().objects.count()


class TestUsuariosCelulas:
    @pytest.mark.parametrize(
        "telefone", ["+32 470 00 00 00", "0470 12 34 56", "0032470123456", "+55 (11) 91234-5678"]
    )
    def test_o_telefone_e_texto_como_foi_cadastrado(self, gerente, telefone):
        _pessoa("Fone", "fone@exemplo.be", phone=telefone)

        folha = _planilha(_exportar_usuarios(gerente, "contatos", q="fone@"))
        celula = folha.cell(row=2, column=3)

        assert celula.value == telefone
        assert celula.data_type == "s"
        assert celula.number_format == "@"

    def test_email_e_nome_sao_texto(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "contatos", q="ana@"))

        assert folha.cell(row=2, column=1).data_type == "s"
        assert folha.cell(row=2, column=2).data_type == "s"

    def test_um_nome_que_parece_formula_continua_texto(self, gerente):
        _pessoa('=HYPERLINK("http://x.test","clique")', "formula@exemplo.be")

        folha = _planilha(_exportar_usuarios(gerente, "completa", q="formula@"))
        celula = folha.cell(row=2, column=1)

        assert celula.value == '=HYPERLINK("http://x.test","clique")'
        assert celula.data_type == "s"

    def test_as_datas_sao_datas(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "completa", q="ana@"))

        assert folha.cell(row=2, column=6).is_date
        assert folha.cell(row=2, column=7).is_date
        assert folha.cell(row=2, column=6).number_format == "DD/MM/YYYY"


class TestUsuariosBuscaEFiltros:
    def test_a_busca_vale(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "completa", q="bruno"))

        assert [linha[1] for linha in _dados(folha)] == ["bruno@exemplo.be"]

    def test_a_busca_pelo_email(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "contatos", q="CARLA@exemplo"))

        assert [linha[1] for linha in _dados(folha)] == ["carla@exemplo.be"]

    def test_o_filtro_de_situacao_vale(self, gerente, pessoas):
        inativos = _planilha(_exportar_usuarios(gerente, "completa", status="inativos"))
        ativos = _planilha(_exportar_usuarios(gerente, "completa", status="ativos"))

        assert [linha[1] for linha in _dados(inativos)] == ["carla@exemplo.be"]
        assert "carla@exemplo.be" not in [linha[1] for linha in _dados(ativos)]
        assert len(_dados(ativos)) == get_user_model().objects.filter(is_active=True).count()

    def test_busca_e_filtro_juntos(self, gerente, pessoas):
        folha = _planilha(
            _exportar_usuarios(gerente, "completa", q="exemplo.be", status="ativos")
        )

        assert sorted(linha[1] for linha in _dados(folha)) == [
            "ana@exemplo.be", "bruno@exemplo.be",
        ]

    def test_sem_resultado_sai_so_o_cabecalho(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "completa", q="ninguem-com-isto"))

        assert _linhas(folha) == [CABECALHO_USUARIOS]

    def test_a_ordem_e_a_da_tela(self, gerente, pessoas):
        folha = _planilha(_exportar_usuarios(gerente, "contatos", q="exemplo.be"))

        assert [linha[0] for linha in _dados(folha)] == [
            "Ana Completa", "Bruno Vazio", "Carla Inativa",
        ]


class TestUsuariosPaginacaoIgnorada:
    """
    57 pessoas batem com a busca; a tela mostra 10 por página e está na
    página 2. A planilha tem de trazer as 57 -- não as 10 da tela.
    """

    @pytest.fixture
    def lote(self, pessoas):
        for numero in range(57):
            _pessoa(f"Lote {numero:02d}", f"lote{numero:02d}@exemplo.be")

    def test_a_tela_mostra_so_uma_pagina(self, gerente, lote):
        resposta = gerente.get(USUARIOS, {"q": "lote", "page_size": "10", "page": "2"})

        assert len(resposta.context["pessoas"]) == 10
        assert resposta.context["total"] == 57

    @pytest.mark.parametrize("tipo", ["completa", "contatos"])
    def test_a_planilha_traz_todos_os_resultados(self, gerente, lote, tipo):
        folha = _planilha(
            _exportar_usuarios(gerente, tipo, q="lote", page="2", page_size="10")
        )

        emails = [linha[1] for linha in _dados(folha)]
        assert len(emails) == 57
        assert emails == [f"lote{numero:02d}@exemplo.be" for numero in range(57)]

    @pytest.mark.parametrize("tamanho", ["10", "20", "30", "50"])
    def test_nenhuma_quantidade_por_pagina_limita(self, gerente, lote, tamanho):
        folha = _planilha(
            _exportar_usuarios(gerente, "contatos", q="lote", page_size=tamanho, page="3")
        )

        assert len(_dados(folha)) == 57

    def test_sem_busca_sai_todo_mundo(self, gerente, lote):
        folha = _planilha(_exportar_usuarios(gerente, "contatos", page_size="10"))

        assert len(_dados(folha)) == get_user_model().objects.count() == 57 + 3 + 1


class TestUsuariosNaTela:
    def test_o_controle_esta_no_cabecalho(self, gerente, pessoas):
        html = gerente.get(USUARIOS).content.decode()
        formulario = _formulario_de_exportacao(html)

        assert "Exportar Excel" in formulario
        assert 'method="post"' in formulario
        assert f'action="{EXPORTAR_USUARIOS}"' in formulario
        assert 'name="csrfmiddlewaretoken"' in formulario
        assert 'name="tipo" value="completa"' in formulario
        assert 'name="tipo" value="contatos"' in formulario
        assert "Tabela completa" in formulario
        assert "Contatos" in formulario

    def test_leva_a_busca_e_o_filtro_mas_nao_a_paginacao(self, gerente, pessoas):
        html = gerente.get(
            USUARIOS, {"q": "exemplo", "status": "ativos", "page_size": "10", "page": "2"}
        ).content.decode()

        assert _escondidos(_formulario_de_exportacao(html)) == {
            "q": "exemplo", "status": "ativos",
        }

    def test_a_tabela_continua_igual(self, gerente, pessoas):
        """A listagem não mudou: as mesmas colunas (com Ações) e a paginação."""
        resposta = gerente.get(USUARIOS, {"page_size": "10"})

        assert resposta.status_code == 200
        assert resposta.context["paginacao"]["tamanho"] == 10
        assert _colunas_da_tabela(resposta.content.decode()) == CABECALHO_USUARIOS + ["Ações"]


# ===========================================================================
# CARTAS
# ===========================================================================


def _carta(pessoa, idioma="fr", convidado="", ida="", volta=""):
    carta = services.start_draft(pessoa, idioma)
    assert carta is not None
    carta.data = {**(carta.data or {}), "guest_name": convidado, "stay_arrival": ida,
                  "stay_departure": volta}
    carta.save(update_fields=["data"])
    return carta


@pytest.fixture
def cartas(modelos_oficiais_prontos, user, other_user):
    """
    `user` (Claire Dubois, com telefone) tem um rascunho em FR com as
    datas da viagem e uma carta FINALIZADA em PT; `other_user` (Rafael
    Costa, sem telefone) tem um rascunho em FR sem datas.
    """
    rascunho = _carta(user, "fr", "Convidado Alfa", "2026-10-01", "2026-10-11")
    finalizada = _carta(user, "pt", "Nome do Rascunho")
    finalizada_em = timezone.make_aware(datetime.datetime(2026, 9, 5, 15, 0))
    Letter.objects.filter(pk=finalizada.pk).update(
        status=Letter.Status.GENERATED,
        finalized_at=finalizada_em,
        snapshot={"data": {"guest_name": "Nome Congelado", "stay_arrival": "2099-05-01",
                           "stay_departure": "2099-05-15"}},
    )
    outra = _carta(other_user, "fr", "Convidada Gama")
    return {
        "rascunho": Letter.objects.get(pk=rascunho.pk),
        "finalizada": Letter.objects.get(pk=finalizada.pk),
        "outra": Letter.objects.get(pk=outra.pk),
    }


def _exportar_cartas(cliente, tipo, **filtros):
    return cliente.post(EXPORTAR_CARTAS, {"tipo": tipo, **filtros})


def _referencias(folha):
    return [linha[0] for linha in _dados(folha)]


class TestCartasPortas:
    def test_anonimo_vai_para_o_login(self, client):
        resposta = client.post(EXPORTAR_CARTAS, {"tipo": "completa"})

        assert resposta.status_code == 302
        assert "/accounts/login/" in resposta.url

    def test_get_nao_exporta(self, supervisora):
        assert supervisora.get(EXPORTAR_CARTAS, {"tipo": "completa"}).status_code == 405

    @pytest.mark.parametrize(
        "permissoes",
        [
            (),
            (BACKOFFICE,),
            (SUPERVISAO,),
            (BACKOFFICE, GERENCIA),
        ],
        ids=["nenhuma", "so-backoffice", "so-supervisao-sem-backoffice", "gerencia-de-usuarios"],
    )
    def test_sem_a_permissao_da_listagem_e_403(self, permissoes):
        cliente = _cliente(*permissoes)

        for tipo in ("completa", "contatos"):
            assert cliente.post(EXPORTAR_CARTAS, {"tipo": tipo}).status_code == 403

    def test_sem_o_token_csrf_e_recusado(self):
        cliente = _cliente(BACKOFFICE, SUPERVISAO, csrf=True)

        assert cliente.post(EXPORTAR_CARTAS, {"tipo": "completa"}).status_code == 403

    @pytest.mark.parametrize("tipo", ["", "tudo", "pdf"])
    def test_tipo_desconhecido_e_400(self, supervisora, tipo):
        assert supervisora.post(EXPORTAR_CARTAS, {"tipo": tipo}).status_code == 400


class TestCartasTabelaCompleta:
    def test_e_um_xlsx_com_o_nome_do_dia(self, supervisora, cartas):
        resposta = _exportar_cartas(supervisora, "completa")

        _planilha(resposta)
        assert resposta["Content-Disposition"] == (
            f'attachment; filename="cartas_tabela_completa_{_hoje()}.xlsx"'
        )

    def test_os_cabecalhos_sao_as_colunas_da_tela_sem_acoes(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa"))

        assert _linhas(folha)[0] == CABECALHO_CARTAS
        assert "Ações" not in _linhas(folha)[0]

    def test_uma_linha_por_carta_na_ordem_da_tela(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa"))

        assert _referencias(folha) == [
            cartas["outra"].reference, cartas["finalizada"].reference,
            cartas["rascunho"].reference,
        ]

    def test_o_rascunho(self, supervisora, cartas):
        carta = cartas["rascunho"]

        linha = _dados(_planilha(_exportar_cartas(supervisora, "completa", q=carta.reference)))[0]

        assert linha[:4] == [carta.reference, "Claire Dubois", "Convidado Alfa", "Français"]
        assert _data(linha[4]) == datetime.date(2026, 10, 1)
        assert _data(linha[5]) == datetime.date(2026, 10, 11)
        assert linha[6] == "Rascunho"
        assert _data(linha[7]) == timezone.localtime(carta.created_at).date()
        assert linha[8] is None

    def test_a_finalizada_mostra_os_dados_congelados(self, supervisora, cartas):
        carta = cartas["finalizada"]

        linha = _dados(_planilha(_exportar_cartas(supervisora, "completa", q=carta.reference)))[0]

        assert linha[2] == "Nome Congelado"
        assert linha[3] == "Português"
        assert _data(linha[4]) == datetime.date(2099, 5, 1)
        assert _data(linha[5]) == datetime.date(2099, 5, 15)
        assert linha[6] == "Finalizada"
        assert _data(linha[8]) == datetime.date(2026, 9, 5)

    def test_dados_ausentes_ficam_vazios(self, supervisora, cartas):
        carta = cartas["outra"]
        Letter.objects.filter(pk=carta.pk).update(data={})

        linha = _dados(_planilha(_exportar_cartas(supervisora, "completa", q=carta.reference)))[0]

        assert linha[1] == "Rafael Costa"
        assert linha[2] is None  # sem convidado
        assert linha[4] is None  # sem data de ida
        assert linha[5] is None  # sem data de volta
        assert linha[8] is None  # não finalizada

    def test_quem_nao_tem_nome_aparece_pelo_email_como_na_tela(self, supervisora, cartas):
        carta = cartas["outra"]
        get_user_model().objects.filter(pk=carta.user_id).update(full_name="")

        linha = _dados(_planilha(_exportar_cartas(supervisora, "completa", q=carta.reference)))[0]

        assert linha[1] == "rafael@exemplo.com"

    def test_cancelada(self, supervisora, cartas):
        carta = cartas["outra"]
        Letter.objects.filter(pk=carta.pk).update(status=Letter.Status.CANCELLED)

        linha = _dados(_planilha(_exportar_cartas(supervisora, "completa", q=carta.reference)))[0]

        assert linha[6] == "Cancelada"

    def test_a_exportacao_nao_mexe_na_carta(self, supervisora, cartas):
        antes = list(Letter.objects.order_by("pk").values())

        _planilha(_exportar_cartas(supervisora, "completa"))
        _planilha(_exportar_cartas(supervisora, "contatos"))

        assert list(Letter.objects.order_by("pk").values()) == antes


class TestCartasContatos:
    def test_so_nome_email_e_telefone_de_quem_criou(self, supervisora, cartas, user, other_user):
        resposta = _exportar_cartas(supervisora, "contatos")
        folha = _planilha(resposta)

        assert resposta["Content-Disposition"] == (
            f'attachment; filename="cartas_contatos_{_hoje()}.xlsx"'
        )
        assert _linhas(folha)[0] == CABECALHO_CONTATOS
        assert folha.max_column == 3
        # Uma linha por carta, na ordem da tela; Rafael não tem telefone.
        assert _dados(folha) == [
            ["Rafael Costa", "rafael@exemplo.com", None],
            ["Claire Dubois", "claire@exemplo.be", "+32 470 00 00 00"],
            ["Claire Dubois", "claire@exemplo.be", "+32 470 00 00 00"],
        ]

    def test_o_telefone_e_texto(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "contatos", q="claire"))

        assert folha.cell(row=2, column=3).value == "+32 470 00 00 00"
        assert folha.cell(row=2, column=3).data_type == "s"

    def test_nada_alem_do_contato(self, supervisora, cartas):
        textos = " | ".join(_textos(_planilha(_exportar_cartas(supervisora, "contatos"))))

        for fora in (cartas["rascunho"].reference, "Convidado Alfa", "Nome Congelado",
                     "Français", "Rascunho", "Finalizada", "Rue des Exemple"):
            assert fora not in textos


class TestCartasBusca:
    def test_pela_referencia(self, supervisora, cartas):
        alvo = cartas["finalizada"]

        folha = _planilha(
            _exportar_cartas(supervisora, "completa", q=alvo.reference[-8:].lower())
        )

        assert _referencias(folha) == [alvo.reference]

    def test_pelo_nome_de_quem_criou(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa", q="claire"))

        assert set(_referencias(folha)) == {
            cartas["rascunho"].reference, cartas["finalizada"].reference,
        }

    def test_pelo_email_de_quem_criou(self, supervisora, cartas, other_user):
        folha = _planilha(_exportar_cartas(supervisora, "contatos", q=other_user.email))

        assert _dados(folha) == [["Rafael Costa", "rafael@exemplo.com", None]]

    def test_pelo_convidado_no_rascunho(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa", q="gama"))

        assert _referencias(folha) == [cartas["outra"].reference]

    def test_pelo_convidado_na_carta_fechada(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa", q="congelado"))

        assert _referencias(folha) == [cartas["finalizada"].reference]

    def test_sem_resultado_sai_so_o_cabecalho(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa", q="nada-bate-com-isto"))

        assert _linhas(folha) == [CABECALHO_CARTAS]


class TestCartasFiltros:
    def test_estado(self, supervisora, cartas):
        rascunhos = _planilha(_exportar_cartas(supervisora, "completa", state="rascunho"))
        finalizadas = _planilha(_exportar_cartas(supervisora, "completa", state="finalizada"))

        assert set(_referencias(rascunhos)) == {
            cartas["rascunho"].reference, cartas["outra"].reference,
        }
        assert _referencias(finalizadas) == [cartas["finalizada"].reference]

    def test_idioma(self, supervisora, cartas):
        folha = _planilha(_exportar_cartas(supervisora, "completa", language="pt"))

        assert _referencias(folha) == [cartas["finalizada"].reference]

    def test_usuario(self, supervisora, cartas, other_user):
        folha = _planilha(_exportar_cartas(supervisora, "completa", user=str(other_user.pk)))

        assert _referencias(folha) == [cartas["outra"].reference]

    def test_periodo_da_viagem(self, supervisora, cartas):
        folha = _planilha(
            _exportar_cartas(supervisora, "completa", **{"from": "2026-09-15", "to": "2026-10-31"})
        )

        assert _referencias(folha) == [cartas["rascunho"].reference]

    def test_filtros_e_busca_juntos(self, supervisora, cartas):
        folha = _planilha(
            _exportar_cartas(supervisora, "contatos", q="claire", language="fr", state="rascunho")
        )

        assert _dados(folha) == [["Claire Dubois", "claire@exemplo.be", "+32 470 00 00 00"]]

    def test_a_exportacao_e_a_tela_concordam(self, supervisora, cartas):
        filtros = {"q": "claire", "state": "rascunho"}

        na_tela = {card.reference for card in supervisora.get(CARTAS, filtros).context["cards"]}
        folha = _planilha(_exportar_cartas(supervisora, "completa", **filtros))

        assert set(_referencias(folha)) == na_tela == {cartas["rascunho"].reference}


class TestCartasPaginacaoIgnorada:
    """
    45 cartas batem com a busca; a tela mostra 10 por página e está na
    página 2. A planilha tem de trazer as 45 -- não as 10 da tela.
    """

    @pytest.fixture
    def lote(self, cartas, user):
        return [_carta(user, "fr", f"Lote {numero:02d}") for numero in range(45)]

    def test_a_tela_mostra_so_uma_pagina(self, supervisora, lote):
        resposta = supervisora.get(CARTAS, {"q": "lote", "page_size": "10", "page": "2"})

        assert len(resposta.context["cards"]) == 10
        assert resposta.context["total"] == 45

    @pytest.mark.parametrize("tipo", ["completa", "contatos"])
    def test_a_planilha_traz_todos_os_resultados(self, supervisora, lote, tipo):
        folha = _planilha(
            _exportar_cartas(supervisora, tipo, q="lote", page="2", page_size="10")
        )

        assert len(_dados(folha)) == 45

    def test_as_cartas_de_todas_as_paginas(self, supervisora, lote):
        pagina_2 = supervisora.get(CARTAS, {"q": "lote", "page_size": "10", "page": "2"})
        folha = _planilha(
            _exportar_cartas(supervisora, "completa", q="lote", page="2", page_size="10")
        )

        exportadas = _referencias(folha)
        assert set(exportadas) == {carta.reference for carta in lote}
        assert {card.reference for card in pagina_2.context["cards"]} < set(exportadas)

    @pytest.mark.parametrize("tamanho", ["10", "20", "30", "50"])
    def test_nenhuma_quantidade_por_pagina_limita(self, supervisora, lote, tamanho):
        folha = _planilha(
            _exportar_cartas(supervisora, "completa", q="lote", state="rascunho",
                             page_size=tamanho, page="3")
        )

        assert len(_dados(folha)) == 45

    def test_sem_busca_saem_todas(self, supervisora, lote):
        folha = _planilha(_exportar_cartas(supervisora, "contatos", page_size="10"))

        assert len(_dados(folha)) == Letter.objects.count() == 45 + 3


class TestCartasNaTela:
    def test_o_controle_esta_no_cabecalho(self, supervisora, cartas):
        html = supervisora.get(CARTAS).content.decode()
        formulario = _formulario_de_exportacao(html)

        assert "Exportar Excel" in formulario
        assert 'method="post"' in formulario
        assert f'action="{EXPORTAR_CARTAS}"' in formulario
        assert 'name="csrfmiddlewaretoken"' in formulario
        assert 'name="tipo" value="completa"' in formulario
        assert 'name="tipo" value="contatos"' in formulario

    def test_leva_busca_e_filtros_mas_nao_a_paginacao(self, supervisora, cartas, user):
        filtros = {"q": "claire", "state": "rascunho", "language": "fr", "user": str(user.pk),
                   "from": "2026-09-01", "to": "2026-12-31"}

        html = supervisora.get(CARTAS, {**filtros, "page_size": "10", "page": "1"}).content.decode()

        assert _escondidos(_formulario_de_exportacao(html)) == filtros

    def test_a_tabela_continua_igual(self, supervisora, cartas):
        """A listagem não mudou: as mesmas colunas (com Ações) e a paginação."""
        resposta = supervisora.get(CARTAS, {"page_size": "10"})

        assert resposta.status_code == 200
        assert resposta.context["paginacao"]["tamanho"] == 10
        assert len(resposta.context["cards"]) == 3
        assert _colunas_da_tabela(resposta.content.decode()) == CABECALHO_CARTAS + ["Ações"]
