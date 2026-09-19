"""
Exportação das tabelas do Backoffice para Excel (.xlsx) -- Usuários e
Cartas.

UM .XLSX DE VERDADE
-------------------
Gerado com openpyxl, não um CSV renomeado: a primeira linha são os
cabeçalhos, depois uma linha por registro. Quem chama entrega as
colunas e as linhas JÁ filtradas; este módulo só sabe escrever.

O TIPO DE CADA CÉLULA
---------------------
  TEXTO  -- gravado sempre como texto, com o formato "@" do Excel: o
            telefone "+32 470 00 00 00" ou "0470..." chega igual, sem
            virar número, sem perder o "+", os zeros nem os espaços. E um
            texto que começa com "=" continua texto -- nunca vira
            fórmula. Um nome de usuário não pode virar `=HYPERLINK(...)`
            vivo na planilha de quem administra.
  DATA   -- valor de data do Excel, exibido como dd/mm/aaaa. Um instante
            (datetime) vira a data dele no fuso do site, a mesma que a
            tabela mostra.
  NUMERO -- o número como número.

`None` ou "" deixa a célula vazia.
"""

import datetime
import io
from dataclasses import dataclass

from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

TEXTO = "texto"
DATA = "data"
NUMERO = "numero"

TIPO_DO_ARQUIVO = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMATO_DE_TEXTO = "@"
FORMATO_DE_DATA = "DD/MM/YYYY"


@dataclass(frozen=True)
class Coluna:
    titulo: str
    tipo: str = TEXTO
    largura: int = 18


def colunas_de_contato():
    """A exportação "Contatos", de Usuários e de Cartas: nome, e-mail, telefone -- nada além."""
    return [
        Coluna(_("Nome"), largura=30),
        Coluna(_("E-mail"), largura=34),
        Coluna(_("Telefone"), largura=20),
    ]


def nome_do_arquivo(prefixo):
    """`<prefixo>_AAAA-MM-DD.xlsx`, com a data de hoje no fuso do site."""
    return f"{prefixo}_{timezone.localdate().isoformat()}.xlsx"


def planilha(colunas, linhas, aba):
    """Os bytes do .xlsx: os cabeçalhos e uma linha por item de `linhas`."""
    livro = Workbook()
    folha = livro.active
    folha.title = str(aba)[:31]

    negrito = Font(bold=True)
    for indice, coluna in enumerate(colunas, start=1):
        celula = folha.cell(row=1, column=indice, value=str(coluna.titulo))
        celula.font = negrito
        folha.column_dimensions[get_column_letter(indice)].width = coluna.largura

    for numero, valores in enumerate(linhas, start=2):
        for indice, (coluna, valor) in enumerate(zip(colunas, valores, strict=True), start=1):
            _gravar(folha.cell(row=numero, column=indice), coluna.tipo, valor)

    # O cabeçalho fica parado quando a planilha rola.
    folha.freeze_panes = "A2"

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def resposta(colunas, linhas, prefixo, aba):
    """A planilha como download: `<prefixo>_AAAA-MM-DD.xlsx`."""
    arquivo = HttpResponse(planilha(colunas, linhas, aba), content_type=TIPO_DO_ARQUIVO)
    arquivo["Content-Disposition"] = f'attachment; filename="{nome_do_arquivo(prefixo)}"'
    # Dado pessoal: nenhum cache guarda uma cópia.
    arquivo["Cache-Control"] = "no-store"
    return arquivo


def _gravar(celula, tipo, valor):
    if valor is None or valor == "":
        return
    if tipo == DATA:
        celula.value = _data(valor)
        celula.number_format = FORMATO_DE_DATA
    elif tipo == NUMERO:
        celula.value = valor
    else:
        # Caracteres de controle invisíveis tornariam o arquivo inválido.
        celula.value = ILLEGAL_CHARACTERS_RE.sub("", str(valor))
        # Texto, sempre: o openpyxl trataria "=..." como fórmula.
        celula.data_type = "s"
        celula.number_format = FORMATO_DE_TEXTO


def _data(valor):
    if isinstance(valor, datetime.datetime):
        if timezone.is_aware(valor):
            return timezone.localtime(valor).date()
        return valor.date()
    return valor
